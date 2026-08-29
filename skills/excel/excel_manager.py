"""
Excel Manager - Core functionality for Excel automation
"""
import xlwings as xw
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
from pathlib import Path
from typing import Optional, List, Dict, Any, Union
import logging

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)


class ExcelManager:
    """Gestor profesional de archivos Excel con xlwings y openpyxl"""
    
    def __init__(self, file_path: Optional[str] = None):
        self.file_path = file_path
        self.wb = None
        self.app = None
        self._openpyxl_wb = None
        
    def open(self, file_path: str, visible: bool = False) -> 'ExcelManager':
        """Abre un archivo Excel existente"""
        self.file_path = file_path
        self.app = xw.App(visible=visible)
        self.wb = self.app.books.open(file_path)
        logger.info(f"Abierto: {file_path}")
        return self
        
    def create(self, file_path: str, visible: bool = False) -> 'ExcelManager':
        """Crea un nuevo archivo Excel"""
        self.file_path = file_path
        self.app = xw.App(visible=visible)
        self.wb = self.app.books.add()
        logger.info(f"Creado: {file_path}")
        return self
        
    def save(self, file_path: Optional[str] = None) -> 'ExcelManager':
        """Guarda el archivo"""
        if file_path:
            self.file_path = file_path
        if self.wb:
            self.wb.save(self.file_path)
            logger.info(f"Guardado: {self.file_path}")
        return self
        
    def close(self) -> None:
        """Cierra el archivo y la aplicación"""
        if self.wb:
            self.wb.close()
        if self.app:
            self.app.quit()
        logger.info("Excel cerrado")
        
    def write_dataframe(self, df: pd.DataFrame, sheet_name: str = "Sheet1", 
                       start_cell: str = "A1", index: bool = False) -> 'ExcelManager':
        """Escribe un DataFrame en una hoja"""
        if sheet_name not in [s.name for s in self.wb.sheets]:
            self.wb.sheets.add(sheet_name)
        sheet = self.wb.sheets[sheet_name]
        sheet.range(start_cell).options(index=index).value = df
        logger.info(f"DataFrame escrito en {sheet_name}!{start_cell}")
        return self
        
    def read_range(self, range_str: str, sheet_name: str = "Sheet1") -> pd.DataFrame:
        """Lee un rango y devuelve DataFrame"""
        sheet = self.wb.sheets[sheet_name]
        data = sheet.range(range_str).value
        return pd.DataFrame(data[1:], columns=data[0])
        
    def format_range(self, range_str: str, sheet_name: str = "Sheet1",
                    bold: bool = False, font_size: int = 11,
                    bg_color: Optional[str] = None,
                    number_format: Optional[str] = None) -> 'ExcelManager':
        """Aplica formato a un rango"""
        sheet = self.wb.sheets[sheet_name]
        rng = sheet.range(range_str)
        
        if bold:
            rng.font.bold = True
        if font_size:
            rng.font.size = font_size
        if bg_color:
            rng.color = bg_color
        if number_format:
            rng.number_format = number_format
            
        return self
        
    def add_chart(self, chart_type: str, data_range: str, 
                  dest_cell: str, sheet_name: str = "Sheet1",
                  title: str = "") -> 'ExcelManager':
        """Agrega un gráfico"""
        sheet = self.wb.sheets[sheet_name]
        
        chart_types = {
            'bar': xw.ChartType.xlColumnClustered,
            'line': xw.ChartType.xlLine,
            'pie': xw.ChartType.xlPie
        }
        
        chart = sheet.charts.add()
        chart.chart_type = chart_types.get(chart_type, xw.ChartType.xlColumnClustered)
        chart.set_source_data(sheet.range(data_range))
        chart.left = sheet.range(dest_cell).left
        chart.top = sheet.range(dest_cell).top
        if title:
            chart.api[1].ChartTitle.Text = title
            
        return self
        
    def apply_template(self, template_name: str, **kwargs) -> 'ExcelManager':
        """Aplica una plantilla predefinida"""
        from .templates import TEMPLATES
        
        if template_name in TEMPLATES:
            TEMPLATES[template_name](self, **kwargs)
            logger.info(f"Plantilla aplicada: {template_name}")
        else:
            logger.warning(f"Plantilla no encontrada: {template_name}")
        return self
        
    def export_to_pdf(self, output_path: Optional[str] = None) -> str:
        """Exporta a PDF"""
        if not output_path:
            output_path = self.file_path.replace('.xlsx', '.pdf')
        self.wb.api.ExportAsFixedFormat(0, output_path)
        logger.info(f"Exportado a PDF: {output_path}")
        return output_path
        
    def __enter__(self):
        return self
        
    def __exit__(self, exc_type, exc_val, exc_tb):
        self.close()
